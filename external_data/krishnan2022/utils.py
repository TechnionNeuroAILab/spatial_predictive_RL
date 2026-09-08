"""Loaders for the Krishnan et al. 2022 (Nat Commun 13:6662) published Source Data.

Source file: MOESM4_source_data.xlsx (see MANIFEST.json for provenance/sha256).
DOI: 10.1038/s41467-022-34465-5

All loaders locate their blocks by scanning for anchor labels (defensive against
blank padding rows / merged cells) and assert expected magnitudes. Every number
returned comes from the xlsx; nothing is hand-typed from the paper text.

Experiment conditions (in-sheet spellings vary; we normalize to snake_case):
    rewarded     -> familiar track, reward delivered at track end
    re_high      -> Reward Extinction, early laps (expectation still high)
    re_low       -> Reward Extinction, late laps (expectation extinguished)
    re_rewarded  -> reward re-introduced
Track: 200 cm virtual linear track, reward zone at the END of the track.
Track regions used throughout the source data:
    start  =   0-50 cm  (farthest from reward)
    middle =  50-150 cm
    end    = 150-200 cm (reward zone / nearest to reward)
"""

from __future__ import annotations

import os
import re
from functools import lru_cache

import pandas as pd

_HERE = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.join(_HERE, "MOESM4_source_data.xlsx")

TRACK_LENGTH_CM = 200.0
REGION_SPANS_CM = {"start": (0.0, 50.0), "middle": (50.0, 150.0), "end": (150.0, 200.0)}


def _norm(x) -> str:
    """Normalize a cell value for label matching (collapse whitespace, lowercase)."""
    if x is None:
        return ""
    return re.sub(r"\s+", " ", str(x)).strip().lower()


@lru_cache(maxsize=None)
def _sheet(name: str):
    """Return worksheet `name` (values only). Cached; workbook loaded once per sheet."""
    import openpyxl

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    if name not in wb.sheetnames:
        raise KeyError(f"sheet {name!r} not in {wb.sheetnames}")
    return wb[name]


def _find_cells(ws, label: str, col=None):
    """All (row, col) whose normalized value equals normalized `label`."""
    target = _norm(label)
    out = []
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and _norm(c.value) == target:
                if col is None or c.column == col:
                    out.append((c.row, c.column))
    return out


def _cell(ws, row: int, col: int):
    return ws.cell(row=row, column=col).value


def _f(x) -> float:
    """Strict float conversion (rejects strings/None so layout drift fails loudly)."""
    if not isinstance(x, (int, float)) or isinstance(x, bool):
        raise ValueError(f"expected numeric cell, got {x!r}")
    return float(x)


# --------------------------------------------------------------------------- #
# Figure 3B: per-animal place-field reliability                                #
# --------------------------------------------------------------------------- #
def load_reliability() -> pd.DataFrame:
    """Per-animal place-field reliability across conditions (sheet 'Figure3').

    Block parsed: 'Figure 3B' -> sub-table titled 'Reliability' (anchor in col A,
    validated by an 'Animal' header directly below), columns
    Animal | Rewarded | Re low | Re-Rewarded, data rows until first non-numeric.

    Returns DataFrame(animal:int, rewarded, re_low, re_rewarded), one row/animal.

    NOTE ON N: the source data contain only n = 3 imaged animals for this panel
    (the paper's behaviour panels, Fig 1F, have n = 12). We assert n >= 3 and
    document the actual n here rather than the >=10 one might expect.
    """
    ws = _sheet("Figure3")
    anchor = None
    for r, c in _find_cells(ws, "Reliability", col=1):
        if _norm(_cell(ws, r + 1, 1)) == "animal":  # header row directly below
            anchor = (r, c)
            break
    if anchor is None:
        raise RuntimeError("Figure3: could not locate 'Reliability' per-animal block")
    hr = anchor[0] + 1
    headers = [_norm(_cell(ws, hr, c)) for c in range(1, 5)]
    assert headers == ["animal", "rewarded", "re low", "re-rewarded"], headers

    rows = []
    r = hr + 1
    while isinstance(_cell(ws, r, 1), (int, float)):
        rows.append(
            dict(
                animal=int(_f(_cell(ws, r, 1))),
                rewarded=_f(_cell(ws, r, 2)),
                re_low=_f(_cell(ws, r, 3)),
                re_rewarded=_f(_cell(ws, r, 4)),
            )
        )
        r += 1
    df = pd.DataFrame(rows)

    assert len(df) >= 3, f"expected >=3 animals, got {len(df)}"  # actual n = 3
    vals = df[["rewarded", "re_low", "re_rewarded"]].to_numpy()
    assert ((vals >= 0.0) & (vals <= 1.0)).all(), "reliability outside [0, 1]"
    return df


# --------------------------------------------------------------------------- #
# Figure 5D: per-axon reward-normalized DA fluorescence ramp                   #
# --------------------------------------------------------------------------- #
_DA_BLOCK_ANCHORS = {
    "rewarded": "Reward Normalized Mean F",
    "re_high": "RE High Normalized Mean F",
    "re_low": "RE low Normalized Mean F",
}


def load_da_ramp(condition: str = "rewarded") -> pd.DataFrame:
    """Per-axon reward-normalized VTA-DA mean fluorescence (sheet 'Figure 5').

    Block parsed: 'Figure 5D' -> sub-table anchored in col B by
    '<Condition> Normalized Mean F' with header Start | Middle | End (cols C-E)
    and rows Axon1..Axon7. `condition` in {'rewarded', 're_high', 're_low'}.

    Track orientation: values RAMP UP from Start to End in the rewarded
    condition, i.e. End of track = reward zone (reward at track end).

    Returns DataFrame(axon:int, start, middle, end), one row per axon (n=7).
    Cross-checks the sheet's own 'Mean' row against the per-axon mean.

    Known source-data quirk: in the 're_low' block, the Middle and End columns
    are identical for every axon (as published; not a parsing error here).
    """
    if condition not in _DA_BLOCK_ANCHORS:
        raise ValueError(f"condition must be one of {sorted(_DA_BLOCK_ANCHORS)}")
    ws = _sheet("Figure 5")
    hits = _find_cells(ws, _DA_BLOCK_ANCHORS[condition], col=2)
    if not hits:
        raise RuntimeError(f"Figure 5: anchor for {condition!r} not found")
    r0 = hits[0][0]
    # Only the first block ('rewarded') carries a Start|Middle|End header row;
    # the RE blocks start their Axon rows directly below the anchor (same cols).
    headers = [_norm(_cell(ws, r0 + 1, c)) for c in (3, 4, 5)]
    if headers == ["start", "middle", "end"]:
        r = r0 + 2
    else:
        assert condition != "rewarded", f"unexpected headers {headers}"
        r = r0 + 1

    rows = []
    while _norm(_cell(ws, r, 2)).startswith("axon"):
        rows.append(
            dict(
                axon=int(re.sub(r"\D", "", str(_cell(ws, r, 2)))),
                start=_f(_cell(ws, r, 3)),
                middle=_f(_cell(ws, r, 4)),
                end=_f(_cell(ws, r, 5)),
            )
        )
        r += 1
    df = pd.DataFrame(rows)

    assert len(df) == 7, f"expected 7 axons, got {len(df)}"
    vals = df[["start", "middle", "end"]].to_numpy()
    assert ((vals > 0.0) & (vals <= 1.0)).all(), "normalized F outside (0, 1]"

    # Cross-check against the sheet's own 'Mean' row (first 'Mean' in col A below block).
    for rr in range(r, r + 6):
        if _norm(_cell(ws, rr, 1)) == "mean":
            for j, col in enumerate(("start", "middle", "end")):
                sheet_mean = _f(_cell(ws, rr, 3 + j))
                assert abs(sheet_mean - df[col].mean()) < 1e-6, (
                    f"{condition}/{col}: sheet mean {sheet_mean} != computed "
                    f"{df[col].mean()}"
                )
            break
    return df


def load_da_ramp_sheet_ci(condition: str = "rewarded") -> dict:
    """The sheet's own '95% CI' row for a Figure 5D block -> {start, middle, end}."""
    ws = _sheet("Figure 5")
    r0 = _find_cells(ws, _DA_BLOCK_ANCHORS[condition], col=2)[0][0]
    for rr in range(r0 + 2, r0 + 20):
        if _norm(_cell(ws, rr, 1)) == "95% ci":
            return {c: _f(_cell(ws, rr, 3 + j)) for j, c in enumerate(("start", "middle", "end"))}
    raise RuntimeError(f"Figure 5: 95% CI row not found for {condition!r}")


# --------------------------------------------------------------------------- #
# Figure 2D: spatial correlation by track region (NOT field density)           #
# --------------------------------------------------------------------------- #
def load_field_density_profiles() -> pd.DataFrame:
    """Per-track-region, per-condition profiles with 95% CIs from sheet 'Figure2'.

    EXACT BLOCK PARSED: 'Figure 2D' (anchor cell 'Figure 2D', region titles
    'Start of track (0-50cm)' / 'Middle of track (50-150 cm)' /
    'End of track (150-200cm)' one row below the anchor, each region a triplet
    Sample Size (cells) | Mean | 95% CI; condition rows Rewarded / RE Low /
    ReRewarded).

    WHAT THIS ACTUALLY IS (decided after inspection): these are SPATIAL
    CORRELATION values (correlation of each cell's spatial activity with its
    rewarded-session template) for place cells grouped by field position on
    the track -- NOT place-field density. The sample sizes are identical
    across conditions (133 / 262 / 207 cells for start/middle/end) because the
    same rewarded-session cells are tracked across conditions (paired t-tests
    in the sheet's Stats block confirm the pairing). Therefore this block
    CANNOT be used for per-condition field-density; per-condition field counts
    by region live in sheet 'Figure3' (Fig 3C sample sizes and Fig 3E
    per-animal percentages) -- see load_field_counts_by_region() and
    load_pf_fraction_by_region().

    Returns long DataFrame(condition, region, region_cm, n_cells,
    mean_spatial_corr, ci95).
    """
    ws = _sheet("Figure2")
    anchor = _find_cells(ws, "Figure 2D")
    if not anchor:
        raise RuntimeError("Figure2: 'Figure 2D' anchor not found")
    r0 = anchor[0][0]

    # Region title cells (row r0+1); each title sits above that region's 'Mean' col.
    region_cols = {}
    for region, label in [
        ("start", "start of track (0-50cm)"),
        ("middle", "middle of track (50-150 cm)"),
        ("end", "end of track (150-200cm)"),
    ]:
        hit = [(r, c) for r, c in _find_cells(ws, label) if r == r0 + 1]
        assert hit, f"Figure 2D: region title {label!r} not found on row {r0 + 1}"
        region_cols[region] = hit[0][1]  # column of 'Mean' for this region

    cond_map = {"rewarded": "rewarded", "re low": "re_low", "rerewarded": "re_rewarded"}
    rows = []
    for rr in range(r0 + 3, r0 + 10):  # condition rows just below the header row
        cond = cond_map.get(_norm(_cell(ws, rr, region_cols["start"] - 2)))
        if cond is None:
            continue
        for region, mean_col in region_cols.items():
            rows.append(
                dict(
                    condition=cond,
                    region=region,
                    region_cm=f"{REGION_SPANS_CM[region][0]:.0f}-{REGION_SPANS_CM[region][1]:.0f}",
                    n_cells=int(_f(_cell(ws, rr, mean_col - 1))),
                    mean_spatial_corr=_f(_cell(ws, rr, mean_col)),
                    ci95=_f(_cell(ws, rr, mean_col + 1)),
                )
            )
    df = pd.DataFrame(rows)

    assert set(df["condition"]) == {"rewarded", "re_low", "re_rewarded"}
    assert len(df) == 9, f"expected 3 conditions x 3 regions, got {len(df)}"
    assert df["mean_spatial_corr"].between(-1, 1).all()
    assert (df["n_cells"] > 50).all()
    # Pairing check: same cells across conditions -> identical n per region.
    assert (df.groupby("region")["n_cells"].nunique() == 1).all(), (
        "n differs across conditions; the paired-cells assumption is wrong"
    )
    return df


# --------------------------------------------------------------------------- #
# Figure 3C: per-condition place-field COUNTS by region (usable as density)    #
# --------------------------------------------------------------------------- #
def load_field_counts_by_region() -> pd.DataFrame:
    """Per-condition place-field counts by track region (sheet 'Figure3', Fig 3C).

    EXACT BLOCK PARSED: 'Figure 3C' -> 'Reliability' sub-table (region titles
    'Start of track (0-50cm)' etc., triplets Sample Size (cells)|Mean|95% CI,
    condition rows Rewarded / RE Low / ReRewarded). Unlike Figure 2D, the
    sample sizes here DIFFER by condition (each condition's own detected place
    fields), so n_fields is a per-condition count of place fields whose
    centroid falls in each region -> usable for density/enrichment.

    Returns long DataFrame(condition, region, region_cm, n_fields,
    mean_reliability, ci95).
    """
    ws = _sheet("Figure3")
    fig3c = [rc for rc in _find_cells(ws, "Figure 3C", col=1)]
    assert fig3c, "Figure3: 'Figure 3C' anchor not found"
    r_fig3c = fig3c[0][0]
    rel = [rc for rc in _find_cells(ws, "Reliability", col=1) if rc[0] > r_fig3c]
    assert rel, "Figure3: Fig 3C 'Reliability' sub-table not found"
    r0 = rel[0][0]

    region_cols = {}
    for region, label in [
        ("start", "start of track (0-50cm)"),
        ("middle", "middle of track (50-150 cm)"),
        ("end", "end of track (150-200cm)"),
    ]:
        hit = [(r, c) for r, c in _find_cells(ws, label) if r == r0 + 1]
        assert hit, f"Fig 3C: region title {label!r} not found on row {r0 + 1}"
        region_cols[region] = hit[0][1]

    cond_map = {"rewarded": "rewarded", "re low": "re_low", "rerewarded": "re_rewarded"}
    rows = []
    for rr in range(r0 + 3, r0 + 10):
        cond = cond_map.get(_norm(_cell(ws, rr, 1)))
        if cond is None:
            continue
        for region, mean_col in region_cols.items():
            rows.append(
                dict(
                    condition=cond,
                    region=region,
                    region_cm=f"{REGION_SPANS_CM[region][0]:.0f}-{REGION_SPANS_CM[region][1]:.0f}",
                    n_fields=int(_f(_cell(ws, rr, mean_col - 1))),
                    mean_reliability=_f(_cell(ws, rr, mean_col)),
                    ci95=_f(_cell(ws, rr, mean_col + 1)),
                )
            )
    df = pd.DataFrame(rows)

    assert len(df) == 9, f"expected 3 conditions x 3 regions, got {len(df)}"
    assert df["mean_reliability"].between(0, 1).all()
    assert (df["n_fields"] > 50).all()
    # Density signal requires n to VARY by condition (unlike Fig 2D).
    assert (df.groupby("region")["n_fields"].nunique() > 1).any()
    return df


# --------------------------------------------------------------------------- #
# Figure 3E: per-animal % of place fields in middle / end of track             #
# --------------------------------------------------------------------------- #
def load_pf_fraction_by_region() -> pd.DataFrame:
    """Per-animal percentage of place fields by track region (sheet 'Figure3', Fig 3E).

    EXACT BLOCKS PARSED: under anchor 'Figure 3E' (col M), three sub-tables
    labelled 'Rewarded' / 'RE low' / 'RE Rewarded', each with header
    Animal | Middle of Track | End of Track and 3 animal rows. The sheet gives
    '% placefields' for middle (50-150 cm) and end (150-200 cm) only;
    pct_start is derived here as 100 - middle - end (start = 0-50 cm).

    Returns DataFrame(condition, animal, pct_middle, pct_end, pct_start).
    """
    ws = _sheet("Figure3")
    fig3e = _find_cells(ws, "Figure 3E")
    assert fig3e, "Figure3: 'Figure 3E' anchor not found"
    r_e, c_e = fig3e[0]

    cond_map = {"rewarded": "rewarded", "re low": "re_low", "re rewarded": "re_rewarded"}
    rows = []
    for rr in range(r_e + 1, r_e + 40):
        cond = cond_map.get(_norm(_cell(ws, rr, c_e)))
        if cond is None:
            continue
        # header row: 'Animal' | 'Middle of Track' | 'End of Track' 1-2 rows below
        hr = None
        for k in (1, 2):
            if _norm(_cell(ws, rr + k, c_e)) == "animal":
                hr = rr + k
                break
        assert hr is not None, f"Fig 3E: 'Animal' header not found under {cond!r}"
        assert _norm(_cell(ws, hr, c_e + 1)) == "middle of track"
        assert _norm(_cell(ws, hr, c_e + 2)) == "end of track"
        r = hr + 1
        while isinstance(_cell(ws, r, c_e), (int, float)):
            mid, end = _f(_cell(ws, r, c_e + 1)), _f(_cell(ws, r, c_e + 2))
            rows.append(
                dict(
                    condition=cond,
                    animal=int(_f(_cell(ws, r, c_e))),
                    pct_middle=mid,
                    pct_end=end,
                    pct_start=100.0 - mid - end,
                )
            )
            r += 1
    df = pd.DataFrame(rows)

    assert set(df["condition"]) == {"rewarded", "re_low", "re_rewarded"}
    assert (df.groupby("condition").size() == 3).all(), "expected 3 animals/condition"
    for col in ("pct_middle", "pct_end", "pct_start"):
        assert df[col].between(0, 100).all(), f"{col} outside [0, 100]"
    return df


if __name__ == "__main__":
    pd.set_option("display.width", 140)
    print("load_reliability():\n", load_reliability(), "\n")
    for cond in ("rewarded", "re_high", "re_low"):
        print(f"load_da_ramp({cond!r}):\n", load_da_ramp(cond), "\n")
    print("load_field_density_profiles():\n", load_field_density_profiles(), "\n")
    print("load_field_counts_by_region():\n", load_field_counts_by_region(), "\n")
    print("load_pf_fraction_by_region():\n", load_pf_fraction_by_region())
